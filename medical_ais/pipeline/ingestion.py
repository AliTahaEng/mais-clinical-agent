"""
Document ingestion — loads raw files (PDF, DOCX, TXT) into text.
Step 1 of the 10-step data pipeline.
"""
from __future__ import annotations

import asyncio
from dataclasses import dataclass, field
from pathlib import Path

import structlog

from medical_ais.core.exceptions import DocumentIngestionError

logger = structlog.get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class RawDocument:
    """Loaded document before any chunking or processing."""
    id: str
    source_path: str
    text: str
    metadata: dict = field(default_factory=dict)


async def load_document(path: Path) -> RawDocument:
    """Load a single file into a RawDocument."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentIngestionError(
            f"Unsupported file type '{suffix}' for {path.name}"
        )

    try:
        if suffix == ".pdf":
            text = await asyncio.to_thread(_read_pdf, path)
        elif suffix == ".docx":
            text = await asyncio.to_thread(_read_docx, path)
        else:
            text = await asyncio.to_thread(path.read_text, encoding="utf-8", errors="replace")
    except DocumentIngestionError:
        raise
    except Exception as exc:
        raise DocumentIngestionError(f"Failed to load '{path.name}': {exc}") from exc

    doc_id = path.stem.lower().replace(" ", "_")
    logger.info("ingestion.loaded", path=str(path), chars=len(text))
    return RawDocument(
        id=doc_id,
        source_path=str(path),
        text=text,
        metadata={"filename": path.name, "extension": suffix},
    )


async def load_directory(directory: str) -> list[RawDocument]:
    """Recursively load all supported documents from *directory*."""
    dir_path = Path(directory)
    if not dir_path.is_dir():
        raise DocumentIngestionError(f"Directory not found: {directory}")

    files = [
        p for p in dir_path.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    logger.info("ingestion.discovered", count=len(files), directory=directory)

    tasks = [load_document(f) for f in files]
    results = await asyncio.gather(*tasks, return_exceptions=True)

    documents: list[RawDocument] = []
    for path, result in zip(files, results):
        if isinstance(result, Exception):
            logger.error("ingestion.failed", path=str(path), error=str(result))
        else:
            documents.append(result)

    return documents


# ── Private file readers ──────────────────────────────────────────────────────

def _read_pdf(path: Path) -> str:
    import pypdf
    reader = pypdf.PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)
